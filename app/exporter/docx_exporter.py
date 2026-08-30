from io import BytesIO

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt


class DOCXExporter:

    def export(
        self,
        report: str,
    ) -> bytes:
        """
        Convert report text into DOCX bytes.

        The document is generated in memory instead
        of being permanently written to disk.
        """

        # Create document
        doc = Document()

        # Process report
        lines = report.split("\n")

        for line in lines:

            line = line.strip()

            if not line:
                continue

            # Main Title
            if line.startswith("# "):

                heading = doc.add_heading(
                    line[2:],
                    level=1,
                )

                heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

            # Section Heading
            elif line.startswith("## "):

                doc.add_heading(
                    line[3:],
                    level=2,
                )

            # Sub Heading
            elif line.startswith("### "):

                doc.add_heading(
                    line[4:],
                    level=3,
                )

            # Bullet Point
            elif line.startswith("* "):

                doc.add_paragraph(
                    line[2:],
                    style="List Bullet",
                )

            elif line.startswith("- "):

                doc.add_paragraph(
                    line[2:],
                    style="List Bullet",
                )

            # Numbered List
            elif (
                len(line) > 2
                and line[0].isdigit()
                and line[1] == "."
            ):

                doc.add_paragraph(
                    line,
                    style="List Number",
                )

            # Normal Paragraph
            else:

                paragraph = doc.add_paragraph()

                run = paragraph.add_run(line)

                run.font.size = Pt(11)

        # Save document to memory
        buffer = BytesIO()

        doc.save(buffer)

        buffer.seek(0)

        return buffer.getvalue()